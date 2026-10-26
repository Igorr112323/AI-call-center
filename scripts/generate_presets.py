"""Generate ready-to-sell demo catalogs (50 SKU each) into presets/.

Run:  python scripts/generate_presets.py
The generated .xlsx files import through the standard price-list parser,
so they exercise the exact production path (synonym headers, extra param
columns, currency-free numeric cells).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PRESETS_DIR = Path(__file__).resolve().parents[1] / "presets"

_HEADER_FILL = PatternFill("solid", fgColor="16191F")
_HEADER_FONT = Font(bold=True, color="E6E8EC")
_CENTER = Alignment(horizontal="center", vertical="center")


def _write_sheet(path: Path, headers: list[str], rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс-лист"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
    for row in rows:
        ws.append(row)
    widths = [16, 52, 16, 14, 10, 14, 14, 14]
    for idx, width in enumerate(widths[:len(headers)], start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ actuators
def actuators() -> tuple[str, list[str], list[list]]:
    headers = ["Артикул", "Наименование", "Категория", "Напряжение, В",
               "Ход штока, мм", "Усилие, Н", "Цена", "Остаток"]
    voltages = (12, 24, 220)
    strokes = (50, 100, 150, 200, 300, 400, 500, 600, 800, 1000)
    rows = []
    for i in range(50):
        voltage = voltages[i % len(voltages)]
        stroke = strokes[(i * 3) % len(strokes)]
        force = 500 + ((i * 350) % 5501)          # 500..6000 Н
        price = 3900 + (stroke * 12) + (force // 20) + voltage * 25
        price = min(48000, round(price / 50) * 50)
        stock = 4 + (i * 7) % 40
        sku = f"LA-{voltage}-{stroke}-{force}"
        name = (f"Актуатор линейный {voltage}В, ход {stroke}мм, "
                f"усилие {force}Н, IP65")
        rows.append([sku, name, "Линейные актуаторы", voltage, stroke,
                     force, price, stock])
    return "actuators_catalog.xlsx", headers, rows


# ------------------------------------------------------------------ RVD hoses
def rvd_hoses() -> tuple[str, list[str], list[list]]:
    headers = ["Артикул", "Наименование", "Категория", "Тип рукава",
               "DN, мм", "Фитинг", "Цена", "Остаток"]
    types = ("1SN", "2SN", "4SP")
    diameters = (6, 8, 10, 12, 16, 20, 25)
    fittings = (("DK", "M14x1.5"), ("DKO", "M16x1.5"), ("DK", "M22x1.5"),
                ("DKO", "M27x2"), ("DK", "M33x2"), ("DKO", "M42x2"))
    rows = []
    for i in range(50):
        hose_type = types[i % len(types)]
        dn = diameters[(i * 2) % len(diameters)]
        fit, thread = fittings[i % len(fittings)]
        length = 500 + ((i * 250) % 2501)          # 0.5..3 м
        price = 580 + dn * 95 + types.index(hose_type) * 420 + length // 10
        price = round(price / 10) * 10
        stock = 8 + (i * 5) % 60
        sku = f"RVD-{hose_type}-DN{dn:02d}-{fit}-{i + 1:02d}"
        name = (f"Рукав высокого давления {hose_type}, DN{dn}, фитинг {fit} "
                f"{thread}, длина {length}мм")
        rows.append([sku, name, "Рукава РВД", hose_type, dn,
                     f"{fit} {thread}", price, stock])
    return "rvd_hydraulics_catalog.xlsx", headers, rows


# ------------------------------------------------------------------ carton boxes
def corrugated_boxes() -> tuple[str, list[str], list[list]]:
    headers = ["Артикул", "Наименование", "Категория", "Марка картона",
               "Тип короба", "Размер ДxШxВ, мм", "Цена", "Остаток"]
    marks = ("Т-21", "Т-23", "Т-24")
    box_types = ("Четырехклапанный", "Самосборный")
    footprints = (
        (200, 150, 100), (300, 200, 150), (400, 300, 200), (500, 350, 250),
        (600, 400, 300), (250, 250, 250), (450, 320, 280), (180, 120, 80),
    )
    rows = []
    for i in range(50):
        mark = marks[i % len(marks)]
        box_type = box_types[i % len(box_types)]
        length, width, height = footprints[(i * 3) % len(footprints)]
        batch = 50 + (i % 10) * 50                  # тираж влияет на цену
        base_price = (
            (length * width * height / 120000)
            * (1.0 + marks.index(mark) * 0.18)
            * (1.15 if box_type == "Самосборный" else 1.0)
            * (1400 / (400 + batch))
        )
        unit_price = max(8.0, round(base_price, 2))
        stock = 100 + (i * 130) % 4000
        sku = f"BOX-{mark.replace('Т-', 'T')}-{i + 1:03d}"
        name = (f"Коробка из гофрокартона {mark}, {box_type.lower()}, "
                f"{length}x{width}x{height}мм, тираж от {batch} шт")
        rows.append([sku, name, "Гофротара", mark, box_type,
                     f"{length}x{width}x{height}", unit_price, stock])
    return "corrugated_boxes_catalog.xlsx", headers, rows


def main() -> None:
    for factory in (actuators, rvd_hoses, corrugated_boxes):
        filename, headers, rows = factory()
        assert len(rows) == 50, f"{filename}: expected 50 SKU, got {len(rows)}"
        _write_sheet(PRESETS_DIR / filename, headers, rows)
        print(f"presets/{filename}: 50 SKU")


if __name__ == "__main__":
    main()
